#!/usr/bin/env python3
"""
One workbook holding all 29 supplementary tables, a sheet each.

  python3 build_combined.py [out.xlsx]

The journal wants the tables one per file, which is what
16_stage_submission_tables.R produces. Reviewers generally prefer one file they
can page through, so this rebuilds the same 29 tables as sheets of a single
workbook, in S1..S29 order, behind a contents sheet.

Nothing is recomputed. Each sheet is a copy of the staged S<N>_Table.xlsx,
same columns, same values, same order, so the two forms cannot drift. The
contents sheet is generated from S_Table_index.csv.

Sheet names are "S<N> <short title>", trimmed to Excel's 31-character limit and
stripped of the characters Excel forbids in a sheet name.
"""
import csv, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import importlib.util as _il, os as _os
_sp = _il.spec_from_file_location("xr", _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "83a_xlsx_reader.py"))
_xr = _il.module_from_spec(_sp); _sp.loader.exec_module(_xr); read_xlsx = _xr.read_xlsx

SRC = ("/mnt/user-data/uploads/SAP-Lipidomics-Database/"
       "final_submission/supplementary_tables")
OUT = sys.argv[1] if len(sys.argv) > 1 else "SoLD_Supplementary_Tables_S1_to_S29.xlsx"

FONT   = "Arial"
HEAD   = Font(name=FONT, size=10, bold=True, color="000000")
BODY   = Font(name=FONT, size=10, color="000000")
TITLE  = Font(name=FONT, size=13, bold=True, color="000000")
FILL   = PatternFill("solid", fgColor="D9D9D9")
THIN   = Side(style="thin", color="000000")
BOX    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BAD    = re.compile(r"[\[\]:*?/\\]")

def sheet_name(n, short, used):
    name = BAD.sub(" ", "S%d %s" % (n, short)).strip()[:31].strip()
    base, k = name, 2
    while name.lower() in used:
        suf = " (%d)" % k; name = base[:31-len(suf)] + suf; k += 1
    used.add(name.lower())
    return name

def write(ws, header, rows):
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.border = HEAD, FILL, BOX
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for r in rows:
        ws.append([r.get(h) for h in header])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY
            cell.alignment = Alignment(vertical="top")
    # width from the longest value, capped so one long cell does not run away
    for c, h in enumerate(header, 1):
        longest = max([len(str(h))] + [len(str(r.get(h, ""))) for r in rows[:400]])
        ws.column_dimensions[get_column_letter(c)].width = min(46, max(9, longest + 2))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), len(rows) + 1)

idx = list(csv.DictReader(open(os.path.join(SRC, "S_Table_index.csv"))))
idx.sort(key=lambda r: int(r["table"][1:]))

wb = Workbook()
toc = wb.active; toc.title = "Contents"
toc["A1"] = "Sorghum Lipid Database (SoLD) — supplementary tables S1 to S29"
toc["A1"].font = TITLE
toc["A2"] = ("One sheet per table, in citation order. Each sheet is identical to the "
             "single-table file of the same number.")
toc["A2"].font = Font(name=FONT, size=10, italic=True, color="000000")
toc.append([]); toc.append([])
toc.append(["Table", "Title", "Sheet", "Rows", "Columns"])
for c in range(1, 6):
    cell = toc.cell(row=5, column=c)
    cell.font, cell.fill, cell.border = HEAD, FILL, BOX

used, total, problems = set(), 0, []
for rec in idx:
    n = int(rec["table"][1:])
    path = os.path.join(SRC, "S%d_Table.xlsx" % n)
    if not os.path.exists(path):
        problems.append("S%d missing" % n); continue
    header, rows = read_xlsx(path)
    name = sheet_name(n, rec["sheet"], used)
    write(wb.create_sheet(name), header, rows)
    toc.append(["S%d" % n, rec["title"].replace("_", " "), name, len(rows), len(header)])
    if int(rec["rows"]) != len(rows) or int(rec["cols"]) != len(header):
        problems.append("S%d: index says %sx%s, file has %dx%d"
                        % (n, rec["rows"], rec["cols"], len(rows), len(header)))
    total += len(rows)
    print("  %-5s %-32s %5d rows x %2d cols" % ("S%d" % n, name, len(rows), len(header)))

for row in toc.iter_rows(min_row=6, max_row=toc.max_row):
    for cell in row:
        cell.font = BODY; cell.border = BOX
for col, w in zip("ABCDE", (9, 42, 34, 9, 10)):
    toc.column_dimensions[col].width = w
toc.append([])
toc.append(["", "%d tables, %d data rows in total" % (len(idx), total)])
toc.cell(row=toc.max_row, column=2).font = Font(name=FONT, size=10, italic=True)
toc.freeze_panes = "A6"

wb.save(OUT)
print("\n%d sheets + contents, %d data rows -> %s (%.1f MB)"
      % (len(wb.sheetnames) - 1, total, OUT, os.path.getsize(OUT) / 1e6))
if problems:
    print("PROBLEMS:"); [print("  ", p) for p in problems]
else:
    print("every sheet matches the row and column counts in S_Table_index.csv")
